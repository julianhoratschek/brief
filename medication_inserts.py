from loaders.patient import Patient
from docx.table import DocxTableRow, DocxTableCellProperties, DocxTableCellWidthProperty
from docx.paragraph import DocxFontProperty, DocxJustificationProperty, DocxSizeProperty, DocxRunProperties
from docx.util import melt

"""
# TODO unused
def get_medication_inserts(patient: Patient):
    # Prepare Table cell widths for column positioning
    c1 = DocxTableCellProperties([DocxTableCellWidthProperty(3686, "dxa")])
    c2 = DocxTableCellProperties([DocxTableCellWidthProperty(1276, "dxa")])
    crest = [DocxTableCellProperties([DocxTableCellWidthProperty(992, "dxa")]),
             DocxTableCellProperties([DocxTableCellWidthProperty(992, "dxa")]),
             DocxTableCellProperties([DocxTableCellWidthProperty(992, "dxa")]),
             DocxTableCellProperties([DocxTableCellWidthProperty(1134, "dxa")])]

    # Base paragraph properties
    ppr = [DocxJustificationProperty(), DocxRunProperties([DocxFontProperty(), DocxSizeProperty(18)])]

    # Helper function for row generation
    def generate_rows(patient_medication) -> list[DocxTableRow]:
        result = []

        for meds in patient_medication:
            row = DocxTableRow()

            # Medication name column
            row.cell(c1).p(ppr).run(meds.name)

            # Medication dosage column
            row.cell(c2).p(ppr).run(f'{meds.amount} {meds.unit}')

            # Time columns, destructure_taken ensures a tuple with 4 elements.
            for cr, t in zip(crest, meds.destructure_taken()):
                row.cell(cr).p(ppr).run(t)

            result.append(row)

        return result

    return {
        'base_medication': melt(generate_rows(patient.current_basis_medication)),
        'other_medication': melt(generate_rows(patient.current_other_medication))
    }

"""